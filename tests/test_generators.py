"""Regression tests for localized templates, allowed quantities and slide images.
Run with: python -m unittest discover -s tests -v
"""
import asyncio
import io
import os
import shutil
import tempfile
import unittest
from pathlib import Path
from unittest.mock import AsyncMock, patch

from docx import Document
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from PIL import Image

from docx_generator import generate_docx, L, _split_into_page_chunks
from pptx_generator import generate_pptx, THEMES, UI, _valid_image_source
from ai_api import (_generate_outline, _is_lang_mixed, _numbered_question_count,
                    _fallback_test_questions)
try:
    from keyboards import option_kb_for, slide_count_kb
    from handlers import creation
    HAS_BOT_RUNTIME = True
except ModuleNotFoundError:
    # The production requirements include aiogram/aiosqlite; generator-only
    # environments can still run DOCX/PPTX and language regression tests.
    option_kb_for = slide_count_kb = creation = None
    HAS_BOT_RUNTIME = False

ROOT = Path(__file__).resolve().parents[1]
LANGS = ("uz", "ru", "en", "tr", "de", "fr")
SERVICES = ("referat", "mustaqil", "kurs", "diplom", "insho", "test", "maqola", "tezis")
COUNTS = (10, 15, 20)
SAMPLE_SLIDES = [
    {"title": "Topic background", "content": ["Historical context and key facts.", "Relevant example.", "Practical significance."]},
    {"title": "Main analysis", "content": ["The first key argument.", "A second point.", "Evidence and implications."]},
    {"title": "Conclusion", "content": ["Summarize the main findings.", "Offer practical recommendations.", "State the result."]},
]
SAMPLE_TEXT = {
    "uz": "Kirish\n\nMavzuning asosiy jihatlari dalillar bilan tahlil qilinadi.\n\nXulosa\nNatijalar umumlashtiriladi.",
    "ru": "Введение\n\nОсновные аспекты темы анализируются на основе фактов.\n\nЗаключение\nРезультаты обобщаются.",
    "en": "Introduction\n\nThe main aspects of the topic are analyzed using evidence.\n\nConclusion\nThe findings are summarized.",
    "tr": "Giriş\n\nKonunun temel yönleri kanıtlarla incelenir.\n\nSonuç\nBulgular özetlenir.",
    "de": "Einleitung\n\nDie wichtigsten Aspekte des Themas werden anhand von Belegen untersucht.\n\nFazit\nDie Ergebnisse werden zusammengefasst.",
    "fr": "Introduction\n\nLes principaux aspects du sujet sont étudiés à partir de faits.\n\nConclusion\nLes résultats sont résumés.",
}
EXTRA_DATA = {
    "ref_subject": "Computer Science", "ref_institution": "Example University",
    "ref_faculty": "Information Technology", "ref_dept": "Computer Science",
    "ref_author": "Test Student, Group 101", "ref_checker": "Test Teacher",
}


class PPTXGeneratorTests(unittest.TestCase):
    def test_allowed_slide_counts_all_designs_and_languages(self):
        for lang in LANGS:
            for design in THEMES:
                for total in COUNTS:
                    with self.subTest(lang=lang, design=design, total=total):
                        stream = generate_pptx("Example topic", "Computer Science", lang, design, total,
                                                SAMPLE_SLIDES, author="Test Student")
                        deck = Presentation(io.BytesIO(stream.getvalue()))
                        self.assertEqual(len(deck.slides), total)
                        last_slide_text = " ".join(shape.text for shape in deck.slides[-1].shapes if hasattr(shape, "text"))
                        self.assertIn(UI[lang]["thanks"], last_slide_text)

    def test_picture_is_embedded_when_image_path_is_provided(self):
        with tempfile.TemporaryDirectory() as tmp:
            image_path = Path(tmp) / "test.jpg"
            Image.new("RGB", (800, 450), (20, 80, 120)).save(image_path, "JPEG")
            stream = generate_pptx("Example topic", "Computer Science", "en", "modern", 10,
                                   SAMPLE_SLIDES, images=[str(image_path)])
            deck = Presentation(io.BytesIO(stream.getvalue()))
            pictures = [shape for slide in deck.slides for shape in slide.shapes if shape.shape_type == MSO_SHAPE_TYPE.PICTURE]
            self.assertGreaterEqual(len(pictures), 1)


    def test_invalid_downloaded_image_is_rejected_and_fallback_can_be_used(self):
        with tempfile.TemporaryDirectory() as tmp:
            bad = Path(tmp) / "not-really-an-image.jpg"
            bad.write_bytes(b"<html>" + b"x" * 3000 + b"</html>")
            self.assertFalse(_valid_image_source(str(bad)))

    def test_picture_is_embedded_even_when_no_external_images_are_available(self):
        stream = generate_pptx("Amir Temur", "History", "uz", "modern", 10,
                               SAMPLE_SLIDES, images=[])
        deck = Presentation(io.BytesIO(stream.getvalue()))
        picture_slides = [slide for slide in deck.slides
                          if any(shape.shape_type == MSO_SHAPE_TYPE.PICTURE for shape in slide.shapes)]
        self.assertGreaterEqual(len(picture_slides), 1)


class DOCXGeneratorTests(unittest.TestCase):
    def test_page_chunks_are_bounded_and_deduplicated(self):
        repeated = "Amir Temur davlat boshqaruvida muhim o‘rin egallaydi. " * 150
        chunks = _split_into_page_chunks(repeated, 5, "Amir Temur", "uz")
        self.assertEqual(len(chunks), 5)
        self.assertTrue(all(len(chunk) <= 1650 for chunk in chunks))
        normalized = [" ".join(chunk.casefold().split()) for chunk in chunks]
        self.assertGreater(len(set(normalized)), 1)

    def test_translation_dictionaries_have_matching_keys(self):
        expected = set(L["uz"])
        for lang in LANGS:
            self.assertEqual(set(L[lang]), expected, msg=f"Missing localization key for {lang}")

    def test_all_services_languages_and_quantity_options_generate(self):
        for lang in LANGS:
            for service in SERVICES:
                for amount in COUNTS:
                    with self.subTest(lang=lang, service=service, amount=amount):
                        stream = generate_docx("Example topic", "Example service", service, lang, amount,
                                               SAMPLE_TEXT[lang], author="Test Student", extra_data=EXTRA_DATA)
                        self.assertGreater(len(stream.getvalue()), 1000)
                        doc = Document(io.BytesIO(stream.getvalue()))
                        all_text = "\n".join(p.text for p in doc.paragraphs)
                        self.assertNotIn("IT sohasi rivojlanmoqda", all_text)
                        self.assertNotIn("Qaysi javob to'g'ri?", all_text)

    def test_supplied_docx_template_page_setup_is_used(self):
        stream = generate_docx("Example topic", "Referat", "referat", "uz", 10,
                               SAMPLE_TEXT["uz"], author="Test Student")
        generated = Document(io.BytesIO(stream.getvalue()))
        template = Document(str(ROOT / "templates" / "referat_IT.docx"))
        self.assertEqual(generated.sections[0].page_width, template.sections[0].page_width)
        self.assertEqual(generated.sections[0].page_height, template.sections[0].page_height)
        self.assertEqual(generated.sections[0].left_margin, template.sections[0].left_margin)


class OptionsAndLanguageTests(unittest.TestCase):
    @unittest.skipUnless(HAS_BOT_RUNTIME, "Install requirements.txt to test Telegram keyboard/handler imports")
    def test_only_10_15_20_choices_are_shown(self):
        pages = [button.callback_data for row in option_kb_for("referat", "uz").inline_keyboard for button in row]
        questions = [button.callback_data for row in option_kb_for("test", "uz").inline_keyboard for button in row]
        slides = [button.callback_data for row in slide_count_kb("uz").inline_keyboard for button in row]
        self.assertEqual([x for x in pages if x.startswith("opt:pages:")], ["opt:pages:10", "opt:pages:15", "opt:pages:20"])
        self.assertEqual([x for x in questions if x.startswith("opt:questions:")], ["opt:questions:10", "opt:questions:15", "opt:questions:20"])
        self.assertEqual([x for x in slides if x.startswith("scnt:")], ["scnt:10", "scnt:15", "scnt:20"])


    def test_language_detector_rejects_wrong_language_output(self):
        samples = {
            "uz": "Amir Temur davlat boshqaruvi tarixida muhim o‘rin egallaydi. U markazlashgan davlat tuzib, ilm-fan va madaniyatni rivojlantirishga e’tibor bergan. Bu jarayonlar mintaqa tarixiga ta’sir ko‘rsatgan.",
            "en": "Amir Timur played an important role in Central Asian history. He built a centralized state and supported architecture, trade, and scholarship. These developments influenced the region for generations.",
            "ru": "Амир Темур сыграл важную роль в истории Центральной Азии. Он создал централизованное государство и поддерживал архитектуру, торговлю и науку.",
            "tr": "Emir Timur, Orta Asya tarihinde önemli bir rol oynadı. Merkezi bir devlet kurdu ve mimariyi, ticareti ve bilimi destekledi.",
            "de": "Amir Timur spielte eine wichtige Rolle in der Geschichte Zentralasiens. Er gründete einen zentralisierten Staat und förderte Architektur, Handel und Wissenschaft.",
            "fr": "Amir Timur a joué un rôle important dans l’histoire de l’Asie centrale. Il a créé un État centralisé et soutenu l’architecture, le commerce et les sciences.",
        }
        for actual_lang, text in samples.items():
            with self.subTest(actual_lang=actual_lang):
                self.assertFalse(_is_lang_mixed(text, actual_lang), f"valid {actual_lang} text rejected")
                for requested_lang in LANGS:
                    if requested_lang != actual_lang:
                        self.assertTrue(_is_lang_mixed(text, requested_lang),
                                        f"{actual_lang} text accepted for {requested_lang}")



    def test_localized_test_fallback_count_and_language(self):
        for lang in LANGS:
            result = _fallback_test_questions("Example topic", 1, 5, lang)
            with self.subTest(lang=lang):
                self.assertEqual(_numbered_question_count(result), 5)
                self.assertFalse(_is_lang_mixed(result, lang))

    def test_test_question_count_parser(self):
        sample = "1. Question one?\nA) One\nB) Two\n2. Question two?\nA) One\nB) Two"
        self.assertEqual(_numbered_question_count(sample), 2)

    def test_body_outline_is_localized(self):
        for lang in LANGS:
            outline = asyncio.run(_generate_outline("referat", "Example topic", 5, lang, {"found": False}))
            self.assertTrue(outline)
            titles = " ".join(item[0] for item in outline)
            if lang == "ru":
                self.assertRegex(titles, r"[А-Яа-я]")
            elif lang == "en":
                self.assertIn("Theoretical", titles)
            elif lang == "de":
                self.assertIn("Grundbegriffe", titles)
            elif lang == "fr":
                self.assertIn("Notions", titles)

    @unittest.skipUnless(HAS_BOT_RUNTIME, "Install requirements.txt to test Telegram keyboard/handler imports")
    def test_image_collection_always_returns_picture_fallback_when_sources_fail(self):
        async def run_test():
            with patch.object(creation, "search_images", new=AsyncMock(return_value=[])), \
                 patch.object(creation, "download_image", new=AsyncMock(return_value=False)):
                paths = await creation._collect_images("Amir Temur", [{"title": "History"}], total=1, lang="uz")
                self.assertTrue(paths)
                self.assertTrue(os.path.isfile(paths[0]))
                with Image.open(paths[0]) as im:
                    self.assertGreaterEqual(im.width, 800)
                return str(Path(paths[0]).parent)
        tmpdir = asyncio.run(run_test())
        shutil.rmtree(tmpdir, ignore_errors=True)


if __name__ == "__main__":
    unittest.main(verbosity=2)
